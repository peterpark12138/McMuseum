# -*- coding: utf-8 -*-
"""产品图总览.xlsx：47 款产品 + 嵌入图片预览"""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

import os, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

BASE = r"D:/WorkbuddyTemp/2026-10-10-21-21-06/McMuseum"
PDIR = os.path.join(BASE, "web", "assets", "products")
THUMBS = os.path.join(BASE, "docs", ".产品图总览.ref", "thumbs")
os.makedirs(THUMBS, exist_ok=True)
OUT = os.path.join(BASE, "docs", "产品图总览.xlsx")

sys.path.insert(0, os.path.join(BASE, "scripts"))
import revival
prods = revival.get_products()

MAPPING = {
    "shake_fries": ("shake_fries.png", "官方菜单图"),
    "purple_sweet_potato_pie": ("purple_sweet_potato_pie.gif", "官方菜单图"),
    "taro_pie": ("taro_pie.gif", "官方菜单图"),
    "big_chicken_steak": ("big_chicken_steak.png", "官方菜单图"),
    "apple_pie": ("apple_pie.gif", "官方菜单图"),
    "shake_fries_fun": ("shake_fries_fun.png", "官方菜单图"),
    "peach_pie": ("peach_pie.gif", "官方菜单图"),
    "mango_pie": ("mango_pie.gif", "官方菜单图"),
    "custard_pie": ("custard_pie.gif", "官方菜单图"),
    "braised_spicy_burger": ("braised_spicy_burger.png", "官方菜单图"),
    "grilled_chicken_burger": ("grilled_chicken_burger.png", "官方菜单图"),
    "kakazu_crispy_chicken": ("../v-crispy.png", "设计稿插画"),
    "matcha_mcfloat": ("../v-mcflurry.png", "设计稿插画"),
    "sakura_mcfloat": ("sakura_mcfloat.jpg", "官方海报·去水印抠图"),
    "garlic_chili_sauce": ("garlic_chili_sauce.jpg", "网络搜集·已白底化"),
    "mc_rice_wrap": ("mc_rice_wrap.jpg", "网络搜集·已白底化"),
    "angus_burger": ("angus_burger.jpg", "网络搜集·已白底化"),
    "cilantro_mcfloat": ("cilantro_mcfloat.png", "用户实拍/替换素材"),
    "youtiao_doujiang": ("youtiao_doujiang.jpg", "网络搜集·已白底化"),
    "treasure_triangle": ("treasure_triangle.png", "用户实拍/替换素材"),
    "strawberry_milkshake": ("strawberry_milkshake.jpg", "网络搜集·已白底化"),
    "milkshake": ("milkshake.jpg", "网络搜集·已白底化"),
    "veggie_cup": ("veggie_cup.jpg", "网络搜集·已白底化"),
    "mccool": ("mccool.jpg", "网络搜集·已白底化"),
    "vcut_wedges": ("vcut_wedges.jpg", "网络搜集·已白底化"),
    "curry_wing": ("curry_wing.jpg", "网络搜集·已白底化"),
    "maple_pancake_burger": ("maple_pancake_burger.jpg", "网络搜集·已白底化"),
}

def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value

XL_HEADER_BG = xl_color("#DA291C")
XL_HEADER_FG = xl_color("#FFFFFF")
XL_OK_BG = xl_color("#E8F7EE")
XL_NO_BG = xl_color("#FFF0F2")
XL_BORDER = xl_color("#E8DFD0")
XL_ZEBRA = xl_color("#FFFBF0")
thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "产品图总览"
wb.properties.title = "产品图总览"

# 列宽：A序号 B图片 C产品名 D产品id E文件名 F来源 G状态 H备注
widths = [6, 15, 24, 24, 30, 18, 14, 40]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 表头（第1行）
headers = ["序号", "图片", "产品名", "产品id", "图片文件名", "来源", "状态", "备注"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True, color=XL_HEADER_FG, size=11)
    cell.fill = PatternFill("solid", fgColor=XL_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:H" + str(len(prods) + 1)

have_n = 0
for i, p in enumerate(prods):
    r = i + 2
    pid, name = p["id"], p["name"]
    zebra = PatternFill("solid", fgColor=XL_ZEBRA) if i % 2 else None
    mapping = MAPPING.get(pid)
    fname = mapping[0] if mapping else ""
    src = mapping[1] if mapping else ""
    img_path = os.path.normpath(os.path.join(PDIR, fname)) if fname else ""
    exists = fname and os.path.exists(img_path)

    ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=3, value=name).font = Font(bold=True, size=11)
    ws.cell(row=r, column=4, value=pid).font = Font(name="Consolas", size=10, color="8A7F72")
    ws.cell(row=r, column=5, value=(fname if exists else f"{pid}.jpg（建议）")).font = Font(name="Consolas", size=10)
    ws.cell(row=r, column=6, value=(src if exists else ""))
    st = ws.cell(row=r, column=7, value=("✅ 有图" if exists else "❌ 缺失"))
    st.alignment = Alignment(horizontal="center", vertical="center")
    note = ""
    if exists:
        have_n += 1
    else:
        note = "找到图片后按本行「图片文件名」列的建议名放入 web/assets/products/，即可接入页面"
        ws.cell(row=r, column=5).font = Font(name="Consolas", size=10, color="C8102E")
    ws.cell(row=r, column=8, value=note).alignment = Alignment(wrapText=True, vertical="center")

    for c in range(1, 9):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER
        if zebra and c != 2:
            cell.fill = zebra
        if c in (3, 5, 6, 8):
            cell.alignment = Alignment(vertical="center", wrapText=(c == 8))
    ws.row_dimensions[r].height = 78

    # 嵌入缩略图
    if exists:
        try:
            thumb = PILImage.open(img_path).convert("RGB")
            if fname.endswith(".gif"):
                thumb.seek(0)
            thumb.thumbnail((90, 90))
            tp = os.path.join(THUMBS, f"{pid}.png")
            thumb.save(tp)
            xi = XLImage(tp)
            xi.width, xi.height = 90, 90
            ws.add_image(xi, f"B{r}")
        except Exception as e:
            ws.cell(row=r, column=2, value=f"图片读取失败 {e}").font = Font(color="C8102E", size=9)
    else:
        ws.cell(row=r, column=2, value="🥫").alignment = Alignment(horizontal="center", vertical="center")

wb.save(OUT)
print(f"saved {OUT} | {len(prods)} products, {have_n} with image")
